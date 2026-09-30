try :

  # Inmporting required modules here
  from datetime import datetime,time,timedelta,date
  from zoneinfo import ZoneInfo
  from openpyxl import load_workbook
  from threading import Thread
  import streamlit as st
  import os,openpyxl,time

  # Initial info of code.
  wb = load_workbook("Library_App.xlsx")

  books_info = wb["books_detailes"]
  membership_info = wb["membership_detailes"]
  library_info = wb["library_info"]

  for i in library_info.iter_rows(min_row = 2 , values_only = True) :
    library_info_row = i

  genres = []
  for i in books_info.iter_rows(min_row = 2, values_only = True) :
    genres.append(i[0])
  genres = set(genres)

  # Helping Functions being used in code with proper docstring.

  def greet_title(num) :
    """This function prints designed title on screen.
    Use this function in this way :
      When you want title on screet just call it."""
    print("="*num)
    print("GURUKOOL LIBRARY".center(num))
    print("="*num)

  def current_date() :
    """This returns current date (according Indian Standard Time) when you call it.
    Use this function in this way :
      When you need current date call it (for some situations store in any variable if required)
    Note : Don't do this --> call it once and store in variable and use that variable when needed.
          It won't give you exact date if your runtime is long. It will give you the same date stored in that variable.
    Tip : Better way is call it every single time when you need, so it give exact date."""
    t = datetime.now(ZoneInfo("Asia/Kolkata"))
    date = t.strftime("%d-%m-%Y")
    return date

  def current_time() :
    """This returns current time when you call it & it's precise and accurate according Indian Standard Time
    Use this function in this way :
      When you need current time call it (for some situations store in any variable if required)
    Note : Don't do this --> call it once and store in variable and use that variable when needed.
          It won't give you exact time if your runtime is long. It will give you the same time stored in that variable.
    Tip : Better way is call it every single time when you need, so it give exact time."""
    t = datetime.now(ZoneInfo("Asia/Kolkata"))
    time = t.strftime("%H:%M:%S")
    return time

  def create_file(name,header) :
    """This function Crates a log sheet in log excel, If it does not exists.
    Use this function in this way :
      1.Call function
      2.Pass two required arguments : 1.Sheet name   2.List of headers
    Note: List of header give your sheet headings and add it to excel file.So give required,accurate headers."""
    if not os.path.exists("Library_App_Log.xlsx") :
      wb1 = openpyxl.Workbook()
      if "Sheet" in wb1.sheetnames :
        del wb1["Sheet"]
      sheet1 = wb1.create_sheet(title = name)
      sheet1.append(header)
      wb1.save("Library_App_Log.xlsx")
      wb1.close()
    wb1 = load_workbook("Library_App_Log.xlsx")
    if name not in wb1.sheetnames :
      sheet1 = wb1.create_sheet(title = name)
      sheet1.append(header)
      wb1.save("Library_App_Log.xlsx")
      wb1.close()

  def update_file(name,info) :
    """This function updates records of your log sheet in log excel, If it already exists.
    Use this function in this way :
      1.Call function
      2.Pass two required arguments : 1.Existing sheet name   2.List of records (according to headers)
    Note : List of records update your sheet's records and add it to excel file.So give accurate records according to headers.
    Tip : Use create_file() function just before using this function."""
    wb2 = load_workbook("Library_App_Log.xlsx")
    sheet2 = wb2[name]
    sheet2.append(info)
    wb2.save("Library_App_Log.xlsx")
    wb2.close()

  def personal_detailes() :
    """This function take personal detaile from user like Name,Mobile Number in input and return list of Name,Mobile Number.
    Use this function in this way :
      1.When you need to take personal detailes from user just call it.
      2.Store in a variable and use that variable as per your need.
      Tip : var[0] gives you Name and var[1] gives you Mobile Number."""
    print("\n Personal detailes :")
    while True :
      name = input("\n    Enter your full name : ").strip().title().split()
      if len(name) == 3 :
        if name[0].isalpha() == False or name[1].isalpha() == False or name[2].isalpha() == False :
          print("\n    Name should contain only alphabets (numbers/special characters is not allowed)!")
        elif len(name[0]) < 2 or len(name[1]) < 2 or len(name[2]) < 2 :
          print("\n    Please enter full name (short form is not allowed)!")
        else :
          name = " ".join(name)
          break
      elif len(name) > 3 or len(name) == 2 :
        print("\n    Please enter appropriate name!")
      else :
        print("\n    Please enter full name!")
    while True :
      mobile = input("    Enter your mobile number : ").strip()
      if mobile.isdigit() == False :
        print("\n    Mobile number should only contain numbers (alphabets/special characters not allowed)!")
      elif mobile.startswith("0") :
        print("\n    Mobile number can't start with 0!")
      elif len(mobile) != 10 :
        print("\n    Mobile number should exact contain 10 digits!")
      else :
        break
    return [name,mobile]

  def display_book_info(books_row) :
    """This fuction display a perticular book's information.
    Use this fuction in this way :
      1.First confirm a book for this.(by choice)
      2.Assign a row of information about that book stored in books_info sheet to variable books_row
      3.Then just call it.
    Note : Don't use this function without assigning the row to variable books_row. Otherwise it won't work and give error!"""
    print(f"\n Name      : {books_row[2]}")
    print(f" Author    : {books_row[3]}")
    print(f" Price     : Rs.{books_row[4]}")
    if books_row[5] == 0 :
      print(f" Available : Currently not available!")
      print(f" About     : {books_row[6]}")
      print(f"\n    Sorry the book '{books_row[2]}' is currently not available, it will be available soon.")
    else :
      if books_row[5] == 1 :
        print(f" Available : {books_row[5]} Book")
      else :
        print(f" Available : {books_row[5]} Books")
      print(f" About     : {books_row[6]}")

  def paymenting() :
    """This fuction displays payment modes and take input of payment mode from user and return that payment mode.
    Use this function in this way :
      1.When bill amount displayed to user in Payment detailes section call it.
      2.Store it in required variable and use in in condition of payment in payment_types.
    Note : A condition is already defined in function of wrong payment mode entered. So apply remaining conditions you want."""
    print(f"\n    Payment we only accept through : {", ".join(library_info_row[6].split(","))}")
    pay = input("    Enter payment mode to proceed further : ").strip().upper()
    if pay not in library_info_row[6].split(",") :
      print("\n    Please enter appropriate payment mode!")
    return pay

  def update_books_decrease(books_row,num = 1) :
    """This fuction decrease copies of book when books are purchased or taken.
    Use this function in this way :
      1.When need of decrease copies of book just call it.
      2.Pass a quantity argument if books purchased more than 1.
    Note : This function takes a argument for quantity but default is 1 (if argument not passed).
    Tip : Call it just after book purchased,issued,taken for read. Pass argument quantity while using it in purchase books."""
    for row_number in range(2 , books_info.max_row + 1) :
      code = books_info.cell(row = row_number , column = 2).value
      if code == books_row[1] :
        books_info.cell(row = row_number , column = 6).value = books_row[5] - num
        wb.save("Library_App.xlsx")
        break

  def update_books_increase(books_row) :
    """This fuction increase copies of book by 1 when book is returned.
    Use this function in this way :
      1.When need of increase copies of book just call it.
    Tip : Call it afuter return book."""
    for row_number in range(2 , books_info.max_row + 1) :
      code = books_info.cell(row = row_number , column = 2).value
      if code == books_row[1] :
        books_info.cell(row = row_number , column = 6).value = books_row[5] + 1
        wb.save("Library_App.xlsx")
        break

  def display_genres() :
    """This function displays all genre names."""
    print("\n Genres :\n")
    print(f"   {", ".join(genres)}")

  def display_books(genres_choice) :
    """This function displays all books with there code and name of chosen genre.
    Use this function in this way :
      1.When you want to show books to user just call it.
      2.Pass a argument containing genre choice."""
    print("\n Books :\n")
    books = []
    for i in books_info.iter_rows(min_row = 2, values_only = True) :
      if genres_choice == i[0] :
        books.append(f"{i[1]} - {i[2]}")
    print(f"   {", ".join(books)}")

  def memberships(msg) :
    """This function displays all memberships with sr.no.
    Use this function in this way :
      1.When you want to show user memberships just call it.
      2.Pass a message in argument which you want to display befor memberships (Eg.Memberships we have :)"""
    print(msg)
    count = 1
    for i in membership_info.iter_rows(min_row = 2 , values_only = True) :
      for j in range(count,membership_info.max_row):
        count = count + 1
        print(f"   {j}.{i[0]}")
        break

  def book_row(book_choice) :
    """This function validates the book user want and returns row containing information of that book.
    Use this function in this way :
      1.When you want required book information to perform further opperations on it just call this function.
      2.Pass a argument which contain book choice taken from user."""
    books_row = None
    for i in books_info.iter_rows(min_row = 2 , values_only = True) :
      if book_choice == i[1] :
        books_row = i
        break
    return books_row

  def remove_expired_stuff() :

    """This function automatically removes expired members when you call it."""
    while True :
      try :
        if os.path.exists("Library_App_Log.xlsx") :
          wbL = load_workbook("Library_App_Log.xlsx")
          if "Registered_Members" in wbL.sheetnames :
            registered_members = wbL["Registered_Members"]
          if "Reading" in wbL.sheetnames :
            reading = wbL["Reading"]
        if "Registered_Members" in wbL.sheetnames :
          if registered_members.max_row > 1 :
            for row_number in range(2 , registered_members.max_row + 1) :
              for i in registered_members.iter_rows(min_row = row_number ,max_row = row_number , values_only = True) :
                a = i[9].split("-")
                if datetime.now(ZoneInfo("Asia/Kolkata")) >= datetime(int(a[2][0:4]),int(a[1]),int(a[0]),int(a[2][6:8]),int(a[2][9:11]),int(a[2][12:14]) , tzinfo=ZoneInfo("Asia/Kolkata")) :
                  registered_members.delete_rows(row_number)
                  wbL.save("Library_App_Log.xlsx")
                  break
        if "Reading" in wbL.sheetnames :
          if reading.max_row > 1 :
            curent_time = datetime.now(ZoneInfo("Asia/Kolkata")).time()
            if curent_time >= library_info_row[5] :
                for row_number in range(2 , reading.max_row + 1) :
                  for i in reading.iter_rows(min_row = row_number ,max_row = row_number , values_only = True) :
                    bookrow = book_row(i[2])
                    update_books_increase(bookrow)
                    reading.delete_rows(row_number)
                    wbL.save("Library_App_Log.xlsx")
                    break
      except :
        pass
      time.sleep(1)

  # Functionalities and Code's working algorithm

  def purchase_book() :
    """This function execute functionality of purchase book."""
    choice = "No"
    if os.path.exists("Library_App_Log.xlsx") :
      wbL = load_workbook("Library_App_Log.xlsx")
      if "Registered_Members" in wbL.sheetnames :
        registered_members = wbL["Registered_Members"]
        if registered_members.max_row > 1 :
          while True :
            choice = input("\n Are you a registered member? : ").strip().title()
            if choice == "Yes" :
              while True :
                # This is to prevent program from crashing in some situations.
                member_row = ["A","A","A","A","A","A","A","A"]
                membership_row = ["A","A","A","A","A","A","A","A"]
                password = 0

                check_member_id = input("\n Enter your Member ID (Press 'Enter' to exit) : ").strip().upper()
                for i in registered_members.iter_rows(min_row = 2 , values_only = True) :
                  if check_member_id == i[0] :
                    member_row = i
                    break
                for i in membership_info.iter_rows(min_row = 2 , values_only = True) :
                  if member_row[3] == i[0] :
                    membership_row = i
                    break
                if check_member_id == "" :
                  break
                elif (len(check_member_id) != 8) or (not check_member_id.startswith("M-")) or (not check_member_id[7:10].isdigit()) or (not check_member_id[2:5].isalpha()) :
                  print("\n Please enter appropriate Member ID.")
                elif check_member_id != member_row[0] :
                  print("\n Member not found!\n Please enter correct Member ID. (May be your membership expired!)")
                elif check_member_id == member_row[0] :
                  while True :
                    password = input(" Enter Password (Press 'Enter' to exit) : ").strip()
                    if password == "" :
                      break
                    elif password != member_row[2] :
                      print("\n Wronge Password!")
                    elif password == member_row[2] :
                      break
                  break
              if password == member_row[2] :
                break
            elif choice == "No" :
              break
            else :
              print("\n Please give your response only in Yes/No!")
    while True :
      # This is to prevent program from crashing in some situations.
      payment = 0
      request_quantity = 0
      purchase_choice = 0

      display_genres()
      genres_choice = input("\n Enter genre you want (Press 'Enter' to exit) : ").strip().title()
      if genres_choice == "" :
        break
      elif genres_choice not in genres :
        print("\n Please enter appropriate genre name!")
      elif genres_choice in genres :
        while True :
          display_books(genres_choice)
          book_choice = input("\n Enter book's code in front of it to proceed further (Press 'Enter' to exit) : ").strip().title()
          books_row = book_row(book_choice)
          if book_choice == "" :
            break
          elif books_row == None :
            print("\n Please enter appropriate book code!")
          elif book_choice == books_row[1] :
            display_book_info(books_row)
            if books_row[5] != 0 :
              while True :
                purchase_choice = input(f"\n Would you like to purchase '{books_row[2]}' book? (Yes/No) : ").strip().title()
                if purchase_choice == "Yes" :
                  if choice == "No" :
                    detailes = personal_detailes()
                  while True :
                    try :
                        quantity = int(input("\n Enter number of books you want to purchase : "))
                        if quantity <= 0 :
                          print("\n    Please enter appropriate quantity (it can't be zero/negative)!")
                        elif quantity > books_row[5] :
                          print(f"\n    Sorry insufficient books!")
                          while True :
                            print(f"\n    Available books : {books_row[5]}")
                            request_quantity = input("    Would you like to buy available books (Yes/No) : ").strip().title()
                            if request_quantity == "Yes" :
                              if choice == "No" :
                                create_file("Customer_Wants_More_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                          "Book Price","Customer Wanted","Available Books","Purchase Status","Date","Time"])
                                update_file("Customer_Wants_More_books",["Not a Member.",detailes[0],detailes[1],books_row[1],books_row[2],books_row[4],
                                          quantity,books_row[5],"Yes",current_date(),current_time()])
                              if choice == "Yes" :
                                create_file("Customer_Wants_More_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                          "Book Price","Customer Wanted","Available Books","Purchase Status","Date","Time"])
                                update_file("Customer_Wants_More_books",[member_row[0],member_row[1],member_row[2],books_row[1],books_row[2],books_row[4],
                                          quantity,books_row[5],"Yes",current_date(),current_time()])
                              break
                            elif request_quantity != "No" :
                              print("\n    Please enter your response in Yes/No only!")
                            elif request_quantity == "No" :
                              if choice == "No" :
                                create_file("Customer_Wants_More_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                            "Book Price","Customer Wanted","Available Books","Purchase Status","Date","Time"])
                                update_file("Customer_Wants_More_books",["Not a Member.",detailes[0],detailes[1],books_row[1],books_row[2],books_row[4],
                                            quantity,books_row[5],"No",current_date(),current_time()])
                              if choice == "Yes" :
                                create_file("Customer_Wants_More_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                            "Book Price","Customer Wanted","Available Books","Purchase Status","Date","Time"])
                                update_file("Customer_Wants_More_books",[member_row[0],member_row[1],member_row[2],books_row[1],books_row[2],books_row[4],
                                            quantity,books_row[5],"No",current_date(),current_time()])
                              break
                          if request_quantity == "No" :
                            break
                        elif quantity <= books_row[5] :
                          print("\n Payment detailes :")
                          if choice == "No" :
                            print(f"\n    You have to pay : Rs.{books_row[4]*quantity}")
                          if choice == "Yes" :
                            print(f"\n    You have to pay : Rs.{int((books_row[4]*quantity)-((books_row[4]*quantity)*(membership_row[6]/100)))}")
                            print(f"    Exclusively {membership_row[6]}% discount only for you!")
                          while True :
                            payment = paymenting()
                            if payment in library_info_row[6].split(",") :
                              print("\n    Payment done successfully!")
                              if choice == "No" :
                                create_file("Purchased_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                            "Book Price","Books Purchased","Discount","Total Amount","Payment Mode","Date","Time"])
                                update_file("Purchased_books",["Not a Member.",detailes[0],detailes[1],books_row[1],books_row[2],books_row[4],
                                            quantity,0,books_row[4]*quantity,payment,current_date(),current_time()])
                              if choice == "Yes" :
                                create_file("Purchased_books",["Member ID","Name","Mobile Number","Book Code","Book Name",
                                            "Book Price","Books Purchased","Discount","Total Amount","Payment Mode","Date","Time"])
                                update_file("Purchased_books",[member_row[0],member_row[1],member_row[2],books_row[1],books_row[2],books_row[4],
                                            quantity,membership_row[6],int((books_row[4]*quantity)-((books_row[4]*quantity)*(membership_row[6]/100))),payment,current_date(),current_time()])
                              update_books_decrease(books_row,quantity)
                              if (books_row[5]-quantity) > 0 and (books_row[5]-quantity) < 6 :
                                create_file("Buy_New_books",["Book Code","Book Name","Book Price",
                                            "Currently Available","Date","Time"])
                                update_file("Buy_New_books",[books_row[1],books_row[2],books_row[4],
                                            books_row[5]-quantity,current_date(),current_time()])
                              elif (books_row[5]-quantity) == 0 :
                                create_file("Hurry_Up_Buy_New_books",["Book Code","Book Name","Book Price","Date","Time"])
                                update_file("Hurry_Up_Buy_New_books",[books_row[1],books_row[2],books_row[4],current_date(),current_time()])
                              break
                          if payment in library_info_row[6].split(",") :
                            break
                    except ValueError :
                      print("\n Please enter only numbers! (alphabets/special characters is not allowed)")
                  if request_quantity == "No" :
                    break
                  elif payment in library_info_row[6].split(",") :
                    break
                elif purchase_choice != "No" :
                  print("\n Please give your response only in Yes/No!")
                elif purchase_choice == "No" :
                  create_file("Rejected_books",["Book Code","Book Name","Author Name",
                              "Book Price","Available Books","Date","Time"])
                  update_file("Rejected_books",[books_row[1],books_row[2],books_row[3],
                              books_row[4],books_row[5],current_date(),current_time()])
                  break
              if purchase_choice == "Yes" :
                if payment in library_info_row[6].split(",") :
                  break
        if purchase_choice == "Yes" :
          if payment in library_info_row[6].split(",") :
            break

  def purchase_membership() :
    """This function execute functionality of purchase membership."""
    if os.path.exists("Library_App_Log.xlsx") :
      wbL = load_workbook("Library_App_Log.xlsx")
      if "Registered_Members" in wbL.sheetnames :
        registered_members = wbL["Registered_Members"]
    else :
      pass
    while True :
      memberships("\n We have following Memberships :")
      try :
        membership_choice = input("\n Which one you want? (Press 'Enter' to exit) : ")
        if membership_choice == "" :
          break
        elif (int(membership_choice) <= (membership_info.max_row - 1)) and (int(membership_choice) > 0) :
          for i in membership_info.iter_rows(min_row = (int(membership_choice) + 1),max_row = (int(membership_choice) + 1),values_only = True) :
            membership_row = i
            break
          print("\n Membership detailes :")
          print(f"\n    Name      : {membership_row[1]}")
          print(f"    Duration  : {membership_row[2]} Days")
          print(f"    Charges   : Rs.{membership_row[3]}")
          if membership_row[0] == "Deluxe" :
            print(f"    Access    : {len((membership_row[4]).split(","))} Books (All)")
          else :
            print(f"    Access    : {len((membership_row[4]).split(","))} Books")
          if membership_row[5] == 1 :
            print(f"    Limit     : {membership_row[5]} Book at a time")
          else :
            print(f"    Limit     : {membership_row[5]} Books at a time")
          print(f"    Discount  : {membership_row[6]}% off on any book purchase")
          while True :
            purchase_choice = input(f"\n Would you like to purchase '{membership_row[1]}'? (Yes/No) : ").strip().title()
            if purchase_choice == "Exit" :
              break
            elif purchase_choice == "Yes" :
              detailes = personal_detailes()
              try :
                for i in registered_members.iter_rows(min_row = 2 , values_only = True) :
                  if (detailes[0] == i[1]) and (detailes[1] == i[2]) :
                    print(f"\n You have already purchased {i[3]} membership.")
                    break
                if (detailes[0] == i[1]) and (detailes[1] == i[2]) :
                  break
              except NameError :
                pass
              print("\n Payment detailes :")
              print(f"\n    You have to pay : Rs.{membership_row[3]}")
              while True :
                payment = paymenting()
                if payment in library_info_row[6].split(",") :
                  print("\n    Payment done successfully!\n\n Member registered successfully.")
                  print(f" Congratulations, {detailes[0].split()[0]}! You are now one of our '{membership_row[0]}' members.")
                  member_id = "M-" + detailes[0].split(" ")[0][:1].upper() + detailes[0].split(" ")[1][:1].upper() + detailes[0].split(" ")[2][:1].upper() + detailes[1][7:10]
                  print(f"\n Your Member ID : {member_id}  &  Your mobile number is it's password.")
                  now = datetime.now(ZoneInfo("Asia/Kolkata"))
                  future_time = now + timedelta(days = membership_row[2])
                  future_time = future_time.strftime("%d-%m-%Y  %H:%M:%S")
                  print(f" Your Membership will expire on {future_time}")
                  print("\n !!! Don't share your Member ID and it's password with anyone !!!")
                  create_file("Registered_Members",["Member ID","Name","Mobile Number","Membership",
                              "Current Issued Books","Charges","Payment Mode","Date","Time","Membership Expiry"])
                  update_file("Registered_Members",[member_id,detailes[0],detailes[1],membership_row[0],
                              0,membership_row[3],payment,current_date(),current_time(),future_time])
                  create_file("Registered_Members_Log",["Member ID","Name","Mobile Number","Membership",
                              "Total Issued Books","Charges","Payment Mode","Date","Time","Membership Expiry"])
                  update_file("Registered_Members_Log",[member_id,detailes[0],detailes[1],membership_row[0],
                              0,membership_row[3],payment,current_date(),current_time(),future_time])
                  break
              break
            elif purchase_choice != "No" :
              print("\n Please give your response only in Yes/No!")
            elif purchase_choice == "No" :
              break
          if purchase_choice == "Yes" :
            break
        else :
          print("\n Please enter valid choice number!")
      except ValueError :
        print("\n Please enter only numbers! (alphabets/special characters is not allowed)")

  def issue_book() :
    """This function execute functionality of issue book."""
    if os.path.exists("Library_App_Log.xlsx") :
      wbL = load_workbook("Library_App_Log.xlsx")
      if "Registered_Members" in wbL.sheetnames :
        registered_members = wbL["Registered_Members"]
      if "Issued_Books" in wbL.sheetnames :
        issued_books = wbL["Issued_Books"]
      if "Registered_Members_Log" in wbL.sheetnames :
        registered_members_log = wbL["Registered_Members_Log"]
      if "Registered_Members" not in wbL.sheetnames :
        print("\n Currently there is no any active member!")
      elif "Registered_Members" in wbL.sheetnames :
        if registered_members.max_row > 1 :
          while True :
            # This is to prevent program from crashing in some situations.
            member_row = ["A","A","A","A","A","A","A","A"]
            membership_row = ["A","A","A","A","A","A","A","A"]
            issue_choice = 0

            check_member_id = input("\n Enter your Member ID (Press 'Enter' to exit) : ").strip().upper()
            for i in registered_members.iter_rows(min_row = 2 , values_only = True) :
              if check_member_id == i[0] :
                member_row = i
                break
            for i in membership_info.iter_rows(min_row = 2 , values_only = True) :
              if member_row[3] == i[0] :
                membership_row = i
                break
            if check_member_id == "" :
              break
            elif check_member_id == member_row[0] :
              while True :
                password = input(" Enter Password (Press 'Enter' to exit) : ").strip()
                if password == "" :
                  break
                elif password != member_row[2] :
                  print("\n Wronge Password!")
                elif password == member_row[2] :
                  print(f"\n Verification Successful!\n    Welcome, {member_row[1].split()[0]}!")
                  if member_row[4] == membership_row[5] :
                    print(f"\n Your Issue Book Limit is reached!\n You can't take more than {membership_row[5]} books at a time.")
                    print("\n To take a book you need to return a book first.")
                  else :
                    while True :
                      display_genres()
                      genres_choice = input("\n Enter genre you want (Press 'Enter' to exit) : ").strip().title()
                      if genres_choice == "" :
                        break
                      elif genres_choice not in genres :
                        print("\n Please enter appropriate genre name!")
                      elif genres_choice in genres :
                        while True :
                          display_books(genres_choice)
                          book_choice = input("\n Enter book's code to proceed further (Press 'Enter' to exit) : ").strip().title()
                          books_row = book_row(book_choice)
                          if book_choice == "" :
                            break
                          elif books_row == None :
                            print("\n Please enter appropriate book code!")
                          elif book_choice == books_row[1] :
                            if book_choice not in membership_row[4].split(",") :
                              print(f"\n You don't have access of '{books_row[2]}' book.")
                            elif book_choice in membership_row[4].split(",") :
                              codes = []
                              if "Issued_Books" in wbL.sheetnames :
                                issued_books_rows = []
                                for i in issued_books.iter_rows(min_row = 2 , values_only = True) :
                                  if member_row[0] == i[0] :
                                    issued_books_rows.append(i)
                                for i in issued_books_rows :
                                  codes.append(i[3])
                              if books_row[1] not in codes :
                                display_book_info(books_row)
                                if books_row[5] != 0 :
                                  while True :
                                    issue_choice = input(f"\n Would you like to take '{books_row[2]}' book? (Yes/No) : ").strip().title()
                                    if issue_choice == "Yes" :
                                      for row_number in range(2, registered_members.max_row + 1) :
                                        id = registered_members.cell(row = row_number , column = 1).value
                                        if id == member_row[0] :
                                          registered_members.cell(row = row_number , column = 5).value = member_row[4] + 1
                                          wbL.save("Library_App_Log.xlsx")
                                          break
                                      for row_number in range(2, registered_members_log.max_row + 1) :
                                        id = registered_members_log.cell(row = row_number , column = 1).value
                                        if id == member_row[0] :
                                          registered_members_log.cell(row = row_number , column = 5).value = member_row[4] + 1
                                          wbL.save("Library_App_Log.xlsx")
                                          break
                                      create_file("Issued_Books",["Member ID","Name","Mobile Number","Book Code","Book Name","Date","Time"])
                                      update_file("Issued_Books",[member_row[0],member_row[1],member_row[2],books_row[1],books_row[2],current_date(),current_time()])
                                      create_file("Issued_Books_Log",["Member ID","Name","Mobile Number","Book Code","Book Name","Date","Time"])
                                      update_file("Issued_Books_Log",[member_row[0],member_row[1],member_row[2],books_row[1],books_row[2],current_date(),current_time()])
                                      update_books_decrease(books_row)
                                      print(f"\n Book '{books_row[2]}' issued successfully.")
                                      print(f"\n Enjoy your book!\n And don't forget to return it before {member_row[9]}.")
                                      break
                                    elif issue_choice != "No" :
                                      print("\n Please give your response only in Yes/No!")
                                    elif issue_choice == "No" :
                                      create_file("Rejected_To_Issue_Book",["Book Code","Book Name","Member ID","Name","Mobile Number","Date","Time"])
                                      update_file("Rejected_To_Issue_Book",[books_row[1],books_row[2],member_row[0],member_row[1],member_row[2],current_date(),current_time()])
                                      break
                                  if issue_choice == "Yes" :
                                    break
                              else :
                                print(f"\n You already have book '{books_row[2]}'.\n You can't take same book again!")
                        if issue_choice == "Yes" :
                          break
                  break
              break
            elif (len(check_member_id) != 8) or (not check_member_id.startswith("M-")) or (not check_member_id[7:10].isdigit()) or (not check_member_id[2:5].isalpha()) :
              print("\n Please enter appropriate Member ID.")
            else :
              print("\n Member not found!\n Please enter correct Member ID. (May be your membership expired!)")
        else :
          print("\n Currently there is no any active member!")
    else :
      print("\n Currently there is no any active member!")

  def read_book() :
    """This function execute functionality of read book."""
    curent_time = datetime.now(ZoneInfo("Asia/Kolkata")).time()
    if library_info_row[4] <= curent_time < library_info_row[5] :
      if os.path.exists("Library_App_Log.xlsx") :
        wbL = load_workbook("Library_App_Log.xlsx")
        if "Reading" in wbL.sheetnames :
          reading = wbL["Reading"]
      else :
        pass
      while True :
        # This is to prevent program from crashing in some situations.
        reading_choice = 0

        display_genres()
        genres_choice = input("\n Enter genre you want (Press 'Enter' to exit) : ").strip().title()
        if genres_choice == "" :
          break
        elif genres_choice not in genres :
          print("\n Please enter appropriate genre name!")
        elif genres_choice in genres :
          while True :
            display_books(genres_choice)
            book_choice = input("\n Enter book's code in front of it to proceed further (Press 'Enter' to exit) : ").strip().title()
            books_row = book_row(book_choice)
            if book_choice == "" :
              break
            elif books_row == None :
              print("\n Please enter appropriate book code!")
            elif book_choice == books_row[1] :
              display_book_info(books_row)
              if books_row[5] != 0 :
                while True :
                  reading_choice = input(f"\n Would you like to take '{books_row[2]}' book for reading? (Yes/No) : ").strip().title()
                  if reading_choice == "Yes" :
                    detailes = personal_detailes()
                    codes = []
                    try :
                      if "Reading" in wbL.sheetnames :
                        reading_row = []
                        for i in reading.iter_rows(min_row = 2 , values_only = True) :
                          if (detailes[0] == i[0]) and (detailes[1] == i[1]) :
                            reading_row.append(i)
                        for i in reading_row :
                          codes.append(i[2])
                    except NameError :
                      pass
                    if books_row[1] not in codes :
                      create_file("Reading",["Name","Mobile Number","Book Code","Book Name","Date","Time"])
                      update_file("Reading",[detailes[0],detailes[1],books_row[1],books_row[2],current_date(),current_time()])
                      create_file("Reading_Log",["Name","Mobile Number","Book Code","Book Name","Date","Time"])
                      update_file("Reading_Log",[detailes[0],detailes[1],books_row[1],books_row[2],current_date(),current_time()])
                      update_books_decrease(books_row)
                      print(f"\n Enjoy your book! And don't forget to return it before {library_info_row[5]}.")
                      break
                    else :
                      print(f"\n You already have book '{books_row[2]}'.\n You can't take same book again!")
                      break
                  elif reading_choice != "No" :
                    print("\n Please give your response only in Yes/No!")
                  elif reading_choice == "No" :
                    break
                if reading_choice == "Yes" :
                  break
          if reading_choice == "Yes" :
            break
    else :
      print(f"\n Sorry, Library is closed! you can't take any book to read.\n The Library is open from {library_info_row[4]} morning to {library_info_row[5]} evening.")

  def return_book() :
    """This function execute functionality of return book."""
    while True :
      if os.path.exists("Library_App_Log.xlsx") :
        wbL = load_workbook("Library_App_Log.xlsx")
        if ("Registered_Members" not in wbL.sheetnames) and ("Reading" not in wbL.sheetnames) :
          print("\n Currently there is no any active member or person reading book!")
          break
        elif ("Registered_Members" in wbL.sheetnames) and ("Reading" in wbL.sheetnames) :
          registered_members = wbL["Registered_Members"]
          reading = wbL["Reading"]
          if (registered_members.max_row == 1) and (reading.max_row == 1) :
            print("\n Currently there is no any active member or person reading book!")
            break
        elif ("Registered_Members" in wbL.sheetnames) and ("Reading" not in wbL.sheetnames) :
          registered_members = wbL["Registered_Members"]
          if registered_members.max_row == 1 :
            print("\n Currently there is no any active member or person reading book!")
            break
        elif ("Registered_Members" not in wbL.sheetnames) and ("Reading" in wbL.sheetnames) :
          reading = wbL["Reading"]
          if reading.max_row == 1 :
            print("\n Currently there is no any active member or person reading book!")
            break
      else :
        print("\n Currently there is no any active member or person reading book!")
        break
      choice = input("\n Are you a registered member? (Press 'Enter' to exit) : ").strip().title()
      if choice == "" :
        break
      elif choice == "Yes" :
        if os.path.exists("Library_App_Log.xlsx") :
          wbL = load_workbook("Library_App_Log.xlsx")
          if "Registered_Members" in wbL.sheetnames :
            registered_members = wbL["Registered_Members"]
          if "Issued_Books" in wbL.sheetnames :
            issued_books = wbL["Issued_Books"]
          if "Registered_Members" not in wbL.sheetnames :
            print("\n Currently there is no any active member!")
            break
          elif "Registered_Members" in wbL.sheetnames :
            if registered_members.max_row > 1 :
              while True :
                # This is to prevent program from crashing in some situations.
                member_row = ["A","A","A","A","A","A","A","A"]

                check_member_id = input("\n Enter your Member ID (Press 'Enter' to exit) : ").strip().upper()
                for i in registered_members.iter_rows(min_row = 2 , values_only = True) :
                  if check_member_id == i[0] :
                    member_row = i
                    break
                if check_member_id == "" :
                  break
                elif (len(check_member_id) != 8) or (not check_member_id.startswith("M-")) or (not check_member_id[7:10].isdigit()) or (not check_member_id[2:5].isalpha()) :
                  print("\n Please enter appropriate Member ID.")
                elif check_member_id != member_row[0] :
                  print("\n Member not found!\n Please enter correct Member ID. (May be your membership expired!)")
                elif check_member_id == member_row[0] :
                  while True :
                    password = input(" Enter Password (Press 'Enter' to exit) : ").strip()
                    if password == "" :
                      break
                    elif password != member_row[2] :
                      print("\n Wronge Password!")
                    elif password == member_row[2] :
                      print(f"\n Verification Successful!\n    Welcome, {member_row[1].split()[0]}!")
                      if "Issued_Books" not in wbL.sheetnames :
                        print("\n You don't have any book to return.")
                        break
                      elif "Issued_Books" in wbL.sheetnames :
                        issued_books_rows = []
                        for i in issued_books.iter_rows(min_row = 2 , values_only = True) :
                          if member_row[0] == i[0] :
                            issued_books_rows.append(i)
                        if len(issued_books_rows) == 0 :
                          print("\n You don't have any book to return.")
                          break
                        else :
                          if len(issued_books_rows) == 1 :
                            print(f"\n You have taken {len(issued_books_rows)} book :")
                          else :
                            print(f"\n You have taken {len(issued_books_rows)} books :")
                          count = 1
                          for i in issued_books_rows :
                            for j in range(count,(len(issued_books_rows)+1)) :
                              count = count + 1
                              print(f"   {j}. {i[3]} - {i[4]}")
                              break
                          while True :
                            return_book_choice = input("\n Enter book's code which you want to return (Press 'Enter' to exit) : ").strip().title()
                            books_row = book_row(return_book_choice)
                            codes = []
                            for i in issued_books_rows :
                              codes.append(i[3])
                            if return_book_choice == "" :
                              break
                            elif books_row == None :
                              print("\n Please enter appropriate book code!")
                            elif return_book_choice not in codes :
                              print("\n You have't taken that book.")
                            elif return_book_choice in codes :
                              for row_number in range(2 , issued_books.max_row + 1) :
                                for i in issued_books.iter_rows(min_row = row_number ,max_row = row_number , values_only = True) :
                                  if (member_row[0] == i[0]) and (return_book_choice == i[3]) :
                                    issued_books.delete_rows(row_number)
                                    wbL.save("Library_App_Log.xlsx")
                                    break
                              for row_number in range(2, registered_members.max_row + 1) :
                                id = registered_members.cell(row = row_number , column = 1).value
                                if id == member_row[0] :
                                  registered_members.cell(row = row_number , column = 5).value = member_row[4] - 1
                                  wbL.save("Library_App_Log.xlsx")
                                  break
                              create_file("Returned_Books",["Name","Mobile Number","Member ID","Book Code","Book Name","Date","Time"])
                              update_file("Returned_Books",[member_row[1],member_row[2],member_row[0],books_row[1],books_row[2],current_date(),current_time()])
                              update_books_increase(books_row)
                              print(f"\n Book '{books_row[2]}' returned successfully.")
                              issued_books_rows = []
                              for i in issued_books.iter_rows(min_row = 2 , values_only = True) :
                                if member_row[0] == i[0] :
                                  issued_books_rows.append(i)
                              if len(issued_books_rows) == 0 :
                                print("\n Now you don't have any book.")
                                break
                              else :
                                if len(issued_books_rows) == 1 :
                                  print(f"\n Now you have {len(issued_books_rows)} book :")
                                else :
                                  print(f"\n Now you have {len(issued_books_rows)} books :")
                                count = 1
                                for i in issued_books_rows :
                                  for j in range(count,(len(issued_books_rows)+1)) :
                                    count = count + 1
                                    print(f"   {j}. {i[3]} - {i[4]}")
                                    break
                                break
                          break
                  break
              if check_member_id == member_row[0] :
                break
            else :
              print("\n Currently there is no any active member!")
              break
        else :
          print("\n Currently there is no any active member!")
          break
      elif choice != "No" :
        print("\n Please give your response only in Yes/No!")
      elif choice == "No" :
        if os.path.exists("Library_App_Log.xlsx") :
          wbL = load_workbook("Library_App_Log.xlsx")
          if "Reading" in wbL.sheetnames :
            reading = wbL["Reading"]
          if "Reading" not in wbL.sheetnames :
            print("\n Currently there is no any person reading book!")
            break
          elif "Reading" in wbL.sheetnames :
            if reading.max_row > 1 :
              detailes = personal_detailes()
              reading_row = []
              for i in reading.iter_rows(min_row = 2 , values_only = True) :
                if (detailes[0] == i[0]) and (detailes[1] == i[1]) :
                  reading_row.append(i)
              if len(reading_row) == 0 :
                print("\n Personal detailes doesn't matched!\n Please enter correct personal detailes.(May be you haven't taken book for reading or returned it)")
                break
              else :
                if len(reading_row) == 1 :
                  print(f"\n Now you have {len(reading_row)} book :")
                else :
                  print(f"\n Now you have {len(reading_row)} books :")
                count = 1
                for i in reading_row :
                  for j in range(count,(len(reading_row)+1)) :
                    count = count + 1
                    print(f"   {j}. {i[2]} - {i[3]}")
                    break
                while True :
                  return_book_choice = input("\n Enter book's code which you want to return (Press 'Enter' to exit) : ").strip().title()
                  books_row = book_row(return_book_choice)
                  codes = []
                  for i in reading_row :
                    codes.append(i[2])
                  if return_book_choice == "" :
                    break
                  elif books_row == None :
                    print("\n Please enter appropriate book code!")
                  elif return_book_choice not in codes :
                    print("\n You have't taken that book.")
                  elif return_book_choice in codes :
                    for row_number in range(2 , reading.max_row + 1) :
                      for i in reading.iter_rows(min_row = row_number ,max_row = row_number , values_only = True) :
                        if (detailes[1] == i[1]) and (return_book_choice == i[2]) :
                          reading.delete_rows(row_number)
                          wbL.save("Library_App_Log.xlsx")
                          break
                    create_file("Returned_Books",["Name","Mobile Number","Member ID","Book Code","Book Name","Date","Time"])
                    update_file("Returned_Books",[detailes[0],detailes[1],"Not a Member",books_row[1],books_row[2],current_date(),current_time()])
                    update_books_increase(books_row)
                    print(f"\n Book '{books_row[2]}' returned successfully.")
                    reading_row = []
                    for i in reading.iter_rows(min_row = 2 , values_only = True) :
                      if (detailes[0] == i[0]) and (detailes[1] == i[1]) :
                        reading_row.append(i)
                    if len(reading_row) == 0 :
                      print("\n Now you don't have any book.")
                      break
                    else :
                      if len(reading_row) == 1 :
                        print(f"\n Now you have {len(reading_row)} book :")
                      else :
                        print(f"\n Now you have {len(reading_row)} books :")
                      count = 1
                      for i in reading_row :
                        for j in range(count,(len(reading_row)+1)) :
                          count = count + 1
                          print(f"   {j}. {i[2]} - {i[3]}")
                          break
                      break
                break
            else :
              print("\n Currently there is no any person reading book!")
              break
        else :
          print("\n Currently there is no any person reading book!")
          break

  def see_issued_books() :
    """This function execute functionality of displaying issued books."""
    if os.path.exists("Library_App_Log.xlsx") :
      wbL = load_workbook("Library_App_Log.xlsx")
      if "Registered_Members" in wbL.sheetnames :
        registered_members = wbL["Registered_Members"]
      if "Issued_Books" in wbL.sheetnames :
        issued_books = wbL["Issued_Books"]
      if "Registered_Members" not in wbL.sheetnames :
        print("\n Currently there is no any active member!")
      elif "Registered_Members" in wbL.sheetnames :
        if registered_members.max_row > 1 :
          while True :
            # This is to prevent program from crashing in some situations.
            member_row = ["A","A","A","A","A","A","A","A"]

            check_member_id = input("\n Enter your Member ID (Press 'Enter' to exit) : ").strip().upper()
            for i in registered_members.iter_rows(min_row = 2 , values_only = True) :
              if check_member_id == i[0] :
                member_row = i
                break
            if check_member_id == "" :
              break
            elif check_member_id == member_row[0] :
              while True :
                password = input(" Enter Password (Press 'Enter' to exit) : ").strip()
                if password == "" :
                  break
                elif password != member_row[2] :
                  print("\n Wronge Password!")
                elif password == member_row[2] :
                  print(f"\n Verification Successful!\n    Welcome, {member_row[1].split()[0]}!")
                  if "Issued_Books" not in wbL.sheetnames :
                    print("\n Currently you haven't taken any book.")
                    break
                  elif "Issued_Books" in wbL.sheetnames :
                    issued_books_rows = []
                    for i in issued_books.iter_rows(min_row = 2 , values_only = True) :
                      if member_row[0] == i[0] :
                        issued_books_rows.append(i)
                    if len(issued_books_rows) == 0 :
                      print("\n Currently you haven't taken any book.")
                      break
                    else :
                      if len(issued_books_rows) == 1 :
                        print(f"\n You have taken {len(issued_books_rows)} book :")
                      else :
                        print(f"\n You have taken {len(issued_books_rows)} books :")
                      count = 1
                      for i in issued_books_rows :
                        for j in range(count,(len(issued_books_rows)+1)) :
                          count = count + 1
                          print(f"   {j}.{i[4]}")
                          break
                      break
              break
            elif (len(check_member_id) != 8) or (not check_member_id.startswith("M-")) or (not check_member_id[7:10].isdigit()) or (not check_member_id[2:5].isalpha()) :
              print("\n Please enter appropriate Member ID.")
            else :
              print("\n Member not found!\n Please enter correct Member ID. (May be your membership expired!)")
        else :
          print("\n Currently there is no any active member!")
    else :
      print("\n Currently there is no any active member!")

  def ratings() :
    """This function execute functionality of taking ratings and reviews from user."""
    if os.path.exists("Library_App_Log.xlsx") :
      wbL = load_workbook("Library_App_Log.xlsx")
      if "Ratings_And_Reviews" in wbL.sheetnames :
        ratings = wbL["Ratings_And_Reviews"]
    else :
      pass
    detailes = personal_detailes()
    while True :
      try :
        for i in ratings.iter_rows(min_row = 2 , values_only = True) :
          if (detailes[0] == i[0]) and (detailes[1] == i[1]) :
            print(f"\n You have already given ratings.")
            break
        if (detailes[0] == i[0]) and (detailes[1] == i[1]) :
          break
      except NameError :
        pass
      try :
        rating = float(input("\n Please rate us from 0 to 5 : "))
        if rating < 0 or rating > 5 :
          print("\n Please give appropriate ratings! (from 0 to 5)")
        else :
          review = input("\n Please leave a review about our Library and Service.(Press 'Enter' to skip) : ").strip().capitalize()
          if review != "" :
            create_file("Ratings_And_Reviews",["Name","Mobile Number","Ratings","Reviews","Date","Time"])
            update_file("Ratings_And_Reviews",[detailes[0],detailes[1],rating,review,current_date(),current_time()])
            break
          else :
            create_file("Ratings_And_Reviews",["Name","Mobile Number","Ratings","Reviews","Date","Time"])
            update_file("Ratings_And_Reviews",[detailes[0],detailes[1],rating,"No Review Leaved.",current_date(),current_time()])
            break
      except ValueError :
        print("\n Please enter only numbers! (alphabets/special characters is not allowed)")

  # Displaying content
  st.set_page_config(
      page_title="Smart Library App",
      page_icon="🔖",
      layout="wide"
  )



  if __name__ == "__main__":
    thread = Thread(target=remove_expired_stuff, daemon=True)
    thread.start()
    greet_title(100)
    print(f"\n Welcome to {library_info_row[0]}!")
    print(f"\n Library Information :")
    print(f"    Name   : {library_info_row[0]}")
    print(f"    Place  : {library_info_row[1]}")
    print(f"    Timing : {library_info_row[4]} to {library_info_row[5]}")
    print(f"    Since  : {library_info_row[2]}")
    print(f"    Staff  : {library_info_row[3]} certified")

    print(f"\n Genres :")
    count = 1
    for i in genres :
      for j in range(count,len(genres)+1):
        count = count + 1
        print(f"    {j}.{i}")
        break

    print(f"\n Books :")
    for i in genres :
      books = []
      for j in books_info.iter_rows(min_row = 2, values_only = True) :
        if j[0] == i :
          books.append(j[2])
      if i == "कादंबरी" :
        print(f"    {i:<17}   -   {", ".join(books)}")
      else :
        print(f"    {i:<18} -   {", ".join(books)}")

    memberships("\n Memberships :")

    # Executing functionalities from here

    while True :
      print()
      greet_title(100)
      print("""\n What do you want further?
   1.Purchase a Book.
   2.Purchase a membership.
   3.Issue a Book.
   4.Read a Book.
   5.Return a Book.
   6.See Issued Books.
   7.Rate Us.""")

      choice_num = input(" Enter your choice number (Press 'Enter' to exit) : ").strip()
      if choice_num == "1" :
        purchase_book()
      elif choice_num == "2" :
        purchase_membership()
      elif choice_num == "3" :
        issue_book()
      elif choice_num == "4" :
        read_book()
      elif choice_num == "5" :
        return_book()
      elif choice_num == "6" :
        see_issued_books()
      elif choice_num == "7" :
        ratings()
      elif choice_num == "" :
        print("\n Thanks for coming, Have a nice day!")
        print(" Hope you liked our service.")
        break
      elif choice_num.isdigit() :
        print("\n Please enter valid choice number!")
      else :
        print("\n Please enter only numbers! (alphabets/special characters is not allowed)")

except FileNotFoundError :
  print(" Please attach excel file first to start the Application. (Containing required information of your library)")
  #print(" Sorry! the Application is not working. We will fix it soon.")

